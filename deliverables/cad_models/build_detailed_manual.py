"""Build the illustrated CAD and model manual for the two Handan designs."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUTPUT = HERE / "和氏璧杯_CAD与模型详细说明书.docx"

FONT_SANS = "Noto Sans SC"
FONT_SERIF = "Noto Serif SC"
DARK = "20343A"
ACCENT = "A64232"
GOLD = "B39055"
PALE = "F2F5F4"
PALE_BLUE = "EAF1F2"
BORDER = "D9D9D9"


def set_run_font(run, name=FONT_SANS, size=10.8, bold=False, color="000000"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=9, color="606A6D")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, bold=True, color="000000")
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        set_run_font(run, size=12.5, bold=True, color="000000")
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(5)
    else:
        set_run_font(run, size=11, bold=True, color="000000")
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.38
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.orphan_control = True
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run, size=10.5)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, size=10.5)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, title in enumerate(headers):
        cell = header.cells[i]
        set_cell_fill(cell, DARK)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        set_run_font(run, size=9.5, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[i])
    for r_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_border(cell)
            set_cell_margins(cell)
            if r_index % 2:
                set_cell_fill(cell, PALE_BLUE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(value))
            set_run_font(run, size=9.3)
            if widths:
                cell.width = Inches(widths[i])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, relative_path, caption, width=6.45):
    path = ROOT / relative_path
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(8)
    c.paragraph_format.keep_with_next = False
    run = c.add_run(caption)
    set_run_font(run, size=9, color="525C60")


def add_page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_SANS
    normal.font.size = Pt(10.8)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_SANS)
    title = doc.styles["Title"]
    title.font.name = FONT_SANS
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_SANS)
    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT_SANS
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_SANS)


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("和氏璧杯设计作品  CAD 与模型详细说明")
    set_run_font(run, size=8.5, color="606A6D")
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("工程说明书  ·  ")
    set_run_font(run, size=9, color="606A6D")
    set_page_number(fp)


def build():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)

    # Cover
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("和氏璧杯设计作品 CAD 与模型详细说明书")
    set_run_font(r, size=26, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("步步生典城市漫游章匣与解围双节机关礼盒")
    set_run_font(r, size=14, color=DARK)
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / "bubushengdian_design/04_renders/hero_open.png"), width=Inches(3.05))
    p.add_run("   ").add_picture(str(ROOT / "jiewei_giftbox_design/04_renders/hero_open.png"), width=Inches(3.05))
    p.paragraph_format.space_after = Pt(18)

    add_table(doc, ["版本", "建模单位", "交付状态", "更新日期"], [["1.0", "mm", "数字几何已验证", "2026-10-01"]], [0.8, 1.0, 2.7, 1.2])
    add_body(doc, "本说明书面向评审、打样厂、三维建模人员和后续迭代者，用于解释两套设计的结构逻辑、尺寸参数、CAD 文件关系、打开方法、制造建议与验证边界。两套模型均已完成 FreeCAD 原生文件、STEP 交换文件和 STL 分件输出。")

    add_page_break(doc)
    add_heading(doc, "阅读指南", 1)
    add_body(doc, "先阅读第 1 章了解交付结论和文件类型；评审或展示人员可重点阅读第 2、3 章；需要修改尺寸、打印或打样的人员应继续阅读第 4、5 章和附录。")
    add_heading(doc, "目录", 2)
    add_numbered(doc, [
        "交付结论与文件类型",
        "步步生典城市漫游章匣",
        "解围双节机关礼盒",
        "CAD 文件使用与参数修改",
        "打印打样和装配建议",
        "验证证据与未验证边界",
        "文件索引和交付校验",
    ])

    add_heading(doc, "1 交付结论与文件类型", 1)
    add_body(doc, "两套作品均已形成可编辑、可交换、可分件打印的数字模型。《步步生典》为 10 个命名实体的章匣系统；《解围》为 13 个命名实体的机关礼盒，并提供闭合、解锁、展开三种 STEP 状态。")
    add_table(doc, ["格式", "用途", "推荐软件", "修改能力"], [
        ["FCStd", "保留命名对象和材料意图的原生 CAD", "FreeCAD 1.1+", "高"],
        ["STEP", "跨 CAD 软件交换及整体尺寸检查", "SolidWorks、Fusion 360、Creo、Rhino", "中"],
        ["STL", "3D 打印、快速样机和网格检查", "Cura、PrusaSlicer、Bambu Studio", "低"],
        ["DXF SVG", "包装刀模、裱纸辅助线和激光切割", "AutoCAD、Illustrator、CorelDRAW", "高"],
        ["Python", "按参数重建模型与导出文件", "FreeCAD Python 控制台", "高"],
    ], [0.8, 2.1, 2.3, 0.8])

    add_heading(doc, "2 步步生典城市漫游章匣", 1)
    add_heading(doc, "2.1 设计定位", 2)
    add_body(doc, "该方案将邯郸成语文化转化为可携带的城市漫游工具。六枚印章与折叠路线图形成“抵达地点、理解故事、完成集章”的使用闭环；中秋与国庆作为旅行和赠礼场景，不取代成语文化主线。")
    add_figure(doc, "bubushengdian_design/04_renders/hero_open.png", "图 1  章匣开启状态，可见六枚印章、印台盒和折叠地图", 6.2)

    add_heading(doc, "2.2 总体尺寸与核心参数", 2)
    add_table(doc, ["参数", "数值", "设计意图"], [
        ["闭合外形", "165 × 120 × 30 mm", "可单手携带，适合礼盒和旅行包"],
        ["外壳壁厚", "1.8 mm", "小批量 ABS PC+ABS 或高强度打印的概念基线"],
        ["章盘单侧间隙", "0.35 mm", "保留数字滑动间隙，实物需根据工艺重标"],
        ["章盘展示行程", "92 mm", "抽出后六枚印章可完整识别和取用"],
        ["单枚章柄", "32 × 32 × 23.6 mm", "兼顾握持、防滚和阵列布局"],
        ["折叠地图代理件", "155 × 110 × 1.2 mm", "表示 480 × 330 mm 展开地图的收纳状态"],
    ], [1.7, 1.7, 3.0])

    add_heading(doc, "2.3 零件构成", 2)
    add_table(doc, ["序号", "零件", "数量", "作用", "建议工艺"], [
        ["1", "保护外壳", "1", "容纳章盘和地图，提供运输保护", "FDM 打印或小批量注塑"],
        ["2", "抽拉章盘", "1", "定位六枚印章和印台盒", "FDM SLA 打印"],
        ["3", "成语印章", "6", "作为城市打卡交互载体", "ABS 章柄加可更换橡胶章面"],
        ["4", "独立印台盒", "1", "隔离水性印泥，降低污染", "标准印台加外嵌套"],
        ["5", "折叠路线图", "1", "路线、故事与集章位置的信息载体", "FSC 纸折页及耐折处理"],
    ], [0.5, 1.3, 0.55, 2.2, 1.6])

    add_figure(doc, "bubushengdian_design/04_renders/exploded.png", "图 2  步步生典爆炸结构图", 6.25)
    add_heading(doc, "2.4 装配与使用逻辑", 2)
    add_numbered(doc, [
        "将六枚章柄按 3 × 2 阵列放入章盘，章面朝下，触觉标记朝上。",
        "印台盒独立放入预留区，避免与地图和章柄直接接触。",
        "折叠地图收纳于外壳上层；使用时先取出地图，再抽出章盘。",
        "抵达对应文化点位后阅读成语故事、完成盖章，最后将所有分件归位。",
    ])
    add_figure(doc, "bubushengdian_design/04_renders/ortho_top.png", "图 3  章盒顶视尺寸与布局参考", 6.1)

    add_heading(doc, "3 解围双节机关礼盒", 1)
    add_heading(doc, "3.1 设计定位与尺寸假设", 2)
    add_body(doc, "该方案以“围魏救赵”的行动次序为机关原型：使用者必须先侧向移动“魏”抽屉，使两枚锁钥片离开“赵”托盘槽位，才能抽出中央托盘。文化故事因此转化为可操作的开盒逻辑，而非表面图案。")
    add_body(doc, "由于用户授权直接完成模型且未提供真实商品，本版本固定为六件中性节礼内装，单件尺寸 60 × 60 × 35 mm。这一假设适合作为月饼、糕点、茶点等小型方盒的空间代理，不代表任何具体品牌商品。")
    add_figure(doc, "jiewei_giftbox_design/04_renders/hero_closed.png", "图 4  解围机关礼盒闭合状态", 6.2)

    add_heading(doc, "3.2 总体参数", 2)
    add_table(doc, ["参数", "数值", "设计意图"], [
        ["闭合外形", "290 × 230 × 75 mm", "容纳 6 件 60 mm 级节礼内装和两级抽拉机构"],
        ["灰板基线", "2.0 mm", "小批量精装礼盒的概念材料厚度"],
        ["包纸代理", "0.18 mm", "用于预估裱纸累积尺寸"],
        ["数字单侧间隙", "0.6 mm", "为纸板滑动预留，实物需按打样复标"],
        ["侧向解锁行程", "42 mm", "使锁钥片离开中央托盘槽位"],
        ["中央展示行程", "145 mm", "展示六件内装和中心月璧元素"],
        ["月璧窗口直径", "108 mm", "形成中秋视觉焦点并强化中央显现"],
    ], [1.8, 1.6, 3.0])

    add_heading(doc, "3.3 零件构成", 2)
    add_table(doc, ["序号", "零件", "数量", "功能", "材料意图"], [
        ["1", "外套筒", "1", "形成封闭外观和月璧窗口", "2 mm 灰板裱纸"],
        ["2", "魏侧抽屉", "1", "承载解锁动作和侧向信息", "灰板精装盒"],
        ["3", "锁钥片", "2", "闭合时限位中央托盘", "高韧纸板或 PP 复合片"],
        ["4", "赵中央托盘", "1", "承载内装与主展示动作", "灰板托盘"],
        ["5", "内衬分格", "1", "定位六件内装并控制碰撞", "EVA 植绒或折叠纸托"],
        ["6", "月璧饰环", "1", "强化中秋视觉中心", "烫金纸卡或薄片"],
        ["7", "中性内装", "6", "体积占位与布局验证", "60 × 60 × 35 mm 代理盒"],
    ], [0.5, 1.25, 0.55, 2.15, 1.6])

    add_heading(doc, "3.4 三状态机关原理", 2)
    add_heading(doc, "状态一  闭合", 3)
    add_body(doc, "两枚锁钥片占据中央托盘槽位，中央托盘无法直接抽出。闭合包络为 290 × 230 × 75 mm。")
    add_figure(doc, "jiewei_giftbox_design/04_renders/hero_closed.png", "图 5  闭合状态", 5.8)
    add_heading(doc, "状态二  解锁", 3)
    add_body(doc, "侧向抽屉向右移动 42 mm，锁钥片随抽屉移动并退出托盘槽位。模型记录的数字释放余量为 1.2 mm。")
    add_figure(doc, "jiewei_giftbox_design/04_renders/hero_unlocked.png", "图 6  侧向抽屉移动后的解锁状态", 5.8)
    add_heading(doc, "状态三  展示", 3)
    add_body(doc, "保持侧向抽屉在解锁位置，将中央托盘抽出 145 mm。展开状态包络为 330 × 360 × 75 mm，可完整显示六件内装。")
    add_figure(doc, "jiewei_giftbox_design/04_renders/hero_open.png", "图 7  中央托盘抽出后的展示状态", 5.8)

    add_heading(doc, "3.5 刀模与图层规则", 2)
    add_body(doc, "刀模以毫米为单位，并提供 DXF 和 SVG 两种可编辑格式。每张图包含 10 × 10 mm 校准方框，打印或导入刀模软件后应先测量该方框，确认没有自动缩放。")
    add_table(doc, ["图层", "颜色", "含义", "制版要求"], [
        ["CUT", "红色", "完全切断线", "刀模或激光切割路径"],
        ["CREASE", "蓝色虚线", "压痕或折叠线", "不得与切线混用"],
        ["GLUE", "灰色", "上胶或禁止上胶提示区", "上机前按图纸备注复核"],
    ], [1.0, 1.1, 2.0, 2.4])
    add_figure(doc, "jiewei_giftbox_design/04_renders/dieline_preview.png", "图 8  刀模文件总览与校准标记", 6.2)
    add_figure(doc, "jiewei_giftbox_design/04_renders/exploded.png", "图 9  解围机关礼盒爆炸结构图", 6.15)

    add_heading(doc, "4 CAD 文件使用与参数修改", 1)
    add_heading(doc, "4.1 直接打开", 2)
    add_numbered(doc, [
        "在 FreeCAD 1.1 或更高版本中打开 FCStd 文件，在树状视图中查看各命名实体。",
        "如使用 SolidWorks、Fusion 360、Creo 或 Rhino，导入 STEP 文件并保持单位为 mm。",
        "打开 STL 时不要再次缩放；首次导入应对照本说明书中的包络尺寸。",
        "DXF 导入后先检查 10 mm 校准框，再分配切线、压痕线和粘结区。",
    ])
    add_heading(doc, "4.2 参数化重建", 2)
    add_body(doc, "源模型由 Python 脚本生成，便于修改尺寸并重复导出。《步步生典》使用 `bubushengdian_design/scripts/design_config.py` 和 `03_cad/build_product.py`；《解围》使用 `jiewei_giftbox_design/scripts/design_config.py`、`03_cad/build_giftbox.py` 与 `03_cad/build_dielines.py`。")
    add_table(doc, ["修改项", "主要参数", "必须联动检查"], [
        ["章匣外形", "width depth height wall", "章盘行程、地图折叠尺寸、内部布局"],
        ["印章规格", "stamp width depth height", "手指握持、分隔尺寸、章面制造"],
        ["礼盒外形", "BOX width depth height", "外套筒、两级抽屉、刀模纸张尺寸"],
        ["真实内装", "SKU width depth height count", "内衬分格、托盘强度、整盒重量与操作力"],
        ["滑动间隙", "clearance release margin", "材料吸湿、裱纸、包边和粘合剂累积误差"],
    ], [1.4, 2.0, 3.2])

    add_heading(doc, "4.3 重建后的必做检查", 2)
    add_bullets(doc, [
        "FreeCAD 文档能够重新打开并完成 recompute，没有空 Shape 或无效实体。",
        "STEP 回读后，《步步生典》为 10 个实体，《解围》每个状态为 13 个实体。",
        "各 STL 文件具有非零包络、非零体积和有效网格，打印前再用切片软件检查自交、悬空和壁厚。",
        "刀模的单位为 mm，图层名和颜色未在软件转换过程中丢失。",
    ])

    add_heading(doc, "5 打印打样和装配建议", 1)
    add_heading(doc, "5.1 步步生典快速样机", 2)
    add_table(doc, ["项目", "建议值", "理由"], [
        ["打印工艺", "FDM 或 SLA", "FDM 用于结构试装，SLA 用于章柄表面和精细展示"],
        ["层高", "0.16-0.20 mm", "在展示质量和打印时间之间取平衡"],
        ["壳体周壁", "至少 4 道", "增强抽拉口和边角耐久性"],
        ["章盘抽拉面", "打印后打磨", "降低层纹导致的滑动阻力"],
        ["印章面", "可更换柔性材料", "不建议用硬质树脂直接代替真实印面"],
    ], [1.3, 1.8, 3.4])

    add_heading(doc, "5.2 解围礼盒纸样", 2)
    add_numbered(doc, [
        "先以 1:1 比例在普通卡纸上输出刀模，检查纸张尺寸和折叠顺序。",
        "再使用 2 mm 灰板制作白样，暂不做全裱纸，测量侧向抽屉和中央托盘的摩擦。",
        "验证锁钥片的插入、释放和回位，至少完成 30 次连续开合。",
        "放入六件代理重量的配重块，检查托盘下挠、单手操作力和抽出后的稳定性。",
        "确定间隙后再加入裱纸、包边和表面工艺，最后重做一次全过程尺寸验证。",
    ])
    add_heading(doc, "5.3 共同验收项目", 2)
    add_table(doc, ["类别", "验收内容", "建议判定"], [
        ["尺寸", "外形、开口、行程、内装间隙", "与数字模型比较并记录实测偏差"],
        ["动作", "抽拉、解锁、回位、防脱", "无卡死、无误解锁、无部件脱落"],
        ["强度", "满载变形、边角、薄弱连接", "使用过程中不影响操作和外观"],
        ["耐久", "连续开合、折页、章面拆装", "实验次数和损伤位置必须留痕"],
        ["运输", "1 m 跌落、振动和内装碰撞", "没有实测前不得声称通过"],
    ], [1.1, 2.9, 2.6])

    add_heading(doc, "6 验证证据与未验证边界", 1)
    add_heading(doc, "6.1 已完成的数字验证", 2)
    add_bullets(doc, [
        "《步步生典》FCStd 回读获得 10 个有效 Shape 对象；STEP 回读获得 10 个实体，包络为 165 × 120 × 30 mm。",
        "《解围》FCStd 回读获得 13 个有效 Shape 对象；闭合、解锁、展开 STEP 均回读为 13 个实体。",
        "《解围》三状态包络分别为 290 × 230 × 75 mm、330 × 230 × 75 mm 和 330 × 360 × 75 mm。",
        "四组刀模的 DXF 单位已设置为 mm，并含有 10 × 10 mm 校准图形。",
        "三个设计目录的自动测试合计 47 项，包括 CAD 报告、刀模、渲染、展板和审计契约。",
    ])
    add_heading(doc, "6.2 不能由数字模型代替的验证", 2)
    add_body(doc, "数字模型可以证明尺寸、实体有效性、数字布局和文件可读性，不能证明真实材料下的摩擦、强度、耐久、运输和安全性。所以在没有实体样机和记录前，作品文案不应使用“已通过跌落”、“已批量生产”、“寿命已验证”或类似表述。")
    add_table(doc, ["未验证项目", "风险", "下一步方法"], [
        ["章盘 0.35 mm 间隙", "打印偏差或表面粗糙导致卡滞", "制作三组不同间隙试片并测力"],
        ["印泥隔离", "泄漏、染色和清洁困难", "用真实印台和纸材做污染测试"],
        ["礼盒 0.6 mm 间隙", "吸湿、裱纸和包边后间隙不足", "白样与裱纸样分别量测"],
        ["锁钥片耐久", "根部撕裂或折痕疲劳", "至少 30 次开合并记录损伤"],
        ["满载托盘", "下挠和操作力过大", "按目标净含量配重后实测"],
    ], [1.8, 2.2, 2.5])

    add_heading(doc, "6.3 参赛使用建议", 2)
    add_bullets(doc, [
        "展板上可表述为“参数化 CAD 模型”、“三状态结构演示”、“已完成 STEP STL DXF 回读检查”。",
        "实体样机完成前，建议使用“概念打样建议”、“数字配合基线”，不使用“量产定型”。",
        "《解围》的中性内装应明确标注为规格假设，不将其包装为已获授权的真实品牌商品。",
        "如有实体白样，应在展板中单独附上实拍照片、尺寸记录和测试次数。",
    ])

    add_heading(doc, "7 文件索引和交付校验", 1)
    add_heading(doc, "7.1 步步生典关键文件", 2)
    add_table(doc, ["文件", "内容"], [
        ["bubushengdian_stamp_kit.FCStd", "FreeCAD 原生装配文档，含 10 个命名实体"],
        ["bubushengdian_stamp_kit.step", "通用 CAD 交换文件"],
        ["meshes/*.stl", "外壳、章盘、六枚印章、印台盒和地图代理件"],
        ["build_product.py", "参数化生成和导出脚本"],
        ["cad_report.json", "包络、体积、实体数、有效性和 SHA-256"],
    ], [2.6, 4.0])
    add_heading(doc, "7.2 解围关键文件", 2)
    add_table(doc, ["文件", "内容"], [
        ["jiewei_giftbox.FCStd", "FreeCAD 原生文档，含 13 个命名实体"],
        ["jiewei_giftbox_closed.step", "机关闭合状态"],
        ["jiewei_giftbox_unlocked.step", "侧向抽屉移出 42 mm 的解锁状态"],
        ["jiewei_giftbox_open.step", "中央托盘再移出 145 mm 的展示状态"],
        ["meshes/*.stl", "外套筒、抽屉、锁钥片、托盘、内衬、月璧和六件内装"],
        ["dielines/*.dxf *.svg", "外套筒、侧抽屉、中央托盘和锁钥片刀模"],
        ["build_giftbox.py build_dielines.py", "模型和刀模的参数化生成脚本"],
    ], [2.7, 3.9])
    add_heading(doc, "7.3 交付包", 2)
    add_table(doc, ["交付包", "内容", "SHA-256"], [
        ["01_bubushengdian_stamp_kit_cad_model.zip", "FCStd STEP STL 渲染图 脚本 说明", "345a7c04...ad5f4af8"],
        ["02_jiewei_giftbox_cad_model.zip", "FCStd 三状态 STEP STL DXF SVG 渲染图 脚本 说明", "c4ce45b5...c563af3f"],
    ], [2.4, 2.8, 1.4])
    add_body(doc, "本说明书采用的 CAD 交付基线提交为 db4ac6d，分支为 codex/handan-double-festival-build。交付包更新后应重新生成 SHA256SUMS.txt，不应继续使用旧校验值。")

    add_heading(doc, "7.4 最终验收清单", 2)
    add_bullets(doc, [
        "能打开两个 FCStd 文件，命名对象数分别为 10 和 13。",
        "能导入四个 STEP 文件，尺寸与本说明书一致。",
        "能逐个导入 STL 分件，没有空文件和零尺寸网格。",
        "DXF 的 10 mm 校准框实测为 10 mm，切线和压痕线可区分。",
        "解压两个 ZIP 时无 CRC 错误，SHA-256 与 SHA256SUMS.txt 一致。",
        "所有关于实体样机、品牌授权和量产性能的表述与真实证据一致。",
    ])

    add_body(doc, "结论：两套方案已达到数字 CAD 交付和概念打样的要求。《步步生典》可直接进入外壳与章盘快速打印；《解围》可直接进入 1:1 纸板白样。经实体样机完成间隙、承重、耐久和运输验证后，才能升级为生产定型文件。", bold_lead="结论：")

    doc.core_properties.title = "和氏璧杯设计作品 CAD 与模型详细说明书"
    doc.core_properties.subject = "步步生典与解围两套方案的工程说明"
    doc.core_properties.author = "项目设计组"
    doc.core_properties.keywords = "CAD, FreeCAD, STEP, STL, DXF, 包装设计, 邯郸"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
