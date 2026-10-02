#!/usr/bin/env python3
"""Build the editable Hanbaofang product-data and competition authorization request."""

from __future__ import annotations

from pathlib import Path
import tempfile
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "jiewei_giftbox_design" / "06_submission" / "邯宝坊产品资料与参赛授权申请书.docx"

BURGUNDY = "6E1F2A"
GOLD = "B9924A"
PALE = "F4F0EA"
MID = "DDD5CB"
TEXT = RGBColor(34, 34, 34)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = MID, size: str = "5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str | None = None, size: float = 8.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial Unicode MS"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None, font_size: float = 8.2):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    hdr = tbl.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        set_cell_text(hdr.cells[idx], text, bold=True, color="FFFFFF", size=8.2)
        set_cell_shading(hdr.cells[idx], BURGUNDY)
        set_cell_border(hdr.cells[idx], BURGUNDY)
        if widths:
            hdr.cells[idx].width = Inches(widths[idx])
    for r_index, row in enumerate(rows):
        cells = tbl.add_row().cells
        for idx, text in enumerate(row):
            set_cell_text(cells[idx], text, size=font_size)
            set_cell_border(cells[idx])
            if r_index % 2 == 1:
                set_cell_shading(cells[idx], PALE)
            if widths:
                cells[idx].width = Inches(widths[idx])
    return tbl


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_status_box(doc: Document, title: str, text: str) -> None:
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(1.35)
    tbl.columns[1].width = Inches(5.5)
    set_cell_text(tbl.cell(0, 0), title, bold=True, color="FFFFFF", size=9)
    set_cell_shading(tbl.cell(0, 0), BURGUNDY)
    set_cell_border(tbl.cell(0, 0), BURGUNDY)
    set_cell_text(tbl.cell(0, 1), text, size=9)
    set_cell_shading(tbl.cell(0, 1), PALE)
    set_cell_border(tbl.cell(0, 1), GOLD)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    paragraph.add_run(" 页")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(4)

    title = doc.styles["Title"]
    title.font.name = "Songti SC"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, color in (("Heading 1", 15, BURGUNDY), ("Heading 2", 11, "3B3B3B")):
        style = doc.styles[name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "2026“和氏璧杯” · 《解围》资料协调文件"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(7.5)
    header.runs[0].font.color.rgb = RGBColor(110, 110, 110)
    add_page_number(section.footer.paragraphs[0])


def normalize_docx_archive(path: Path) -> None:
    """Normalize ZIP entry order and timestamps so rebuilds have a stable hash."""
    fixed_time = (2026, 10, 3, 0, 0, 0)
    with zipfile.ZipFile(path, "r") as source:
        entries = [(name, source.read(name)) for name in sorted(source.namelist())]
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=path.parent) as handle:
        temp_path = Path(handle.name)
    try:
        with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
            for name, data in entries:
                info = zipfile.ZipInfo(name, fixed_time)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o644 << 16
                target.writestr(info, data)
        temp_path.replace(path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def build() -> Path:
    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    p.add_run("邯宝坊产品资料与参赛授权申请书")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run("项目：《解围——邯宝坊双节机关礼盒》")
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(BURGUNDY)
    r.bold = True

    add_status_box(doc, "文件状态", "待邯宝坊有权代表回填、签署并盖章；当前不构成授权，也不证明具体 SKU 已锁定。")
    doc.add_paragraph()
    table(
        doc,
        ["项目", "填写内容"],
        [
            ["申请方/作者", "____________________________"],
            ["院校/单位", "____________________________"],
            ["联系人及手机", "____________________________"],
            ["电子邮箱", "____________________________"],
            ["收件方", "邯郸市邯宝坊特产商贸有限公司 品牌管理/赛事对接负责人"],
            ["申请日期", "______年____月____日"],
        ],
        [1.65, 5.1],
        9,
    )
    add_heading(doc, "申请目的", 1)
    add_body(doc, "本申请用于参加 2026“和氏璧杯”成语之都大学生创意奖赛道一“寻味邯郸·农特产品包装设计”。设计以“围魏救赵”的双向抽拉解锁为结构叙事，拟装入六件由邯宝坊指定的真实在售商品。为避免虚构商品事实、误用品牌资产或以净含量代替结构数据，现申请品牌方提供商品资料并确认有限参赛授权。")
    add_body(doc, "本文件不请求生产或销售许可。未经另行书面同意，申请方不会把邯宝坊 Logo、商品图片或标签信息用于商业销售、广告投放、联名宣传或永久作品集发布。")

    doc.add_page_break()
    add_heading(doc, "一、请品牌方指定六个真实 SKU", 1)
    add_body(doc, "以下仅为公开证据形成的候选品类。品牌方可确认、替换或删减；最终以回填的商品全称、条码和样品为唯一建模依据。")
    table(
        doc,
        ["序号", "候选品类/商品", "公开可核验信息", "品牌方最终指定（商品全称+条码）"],
        [
            ["1", "学步桥小磨芝麻香油", "抽检信息见 240 mL/瓶", "________________________________"],
            ["2", "黄粱梦小米", "公开存在 500 g、2.5 kg 等规格", "________________________________"],
            ["3", "鸡泽辣椒酱", "公开存在多品牌及 180–200 g 瓶装规格", "________________________________"],
            ["4", "魏县 NFC 鸭梨饮品", "公开报道与邯宝坊线上线下销售有关联", "________________________________"],
            ["5", "涉县核桃/核桃制品", "邯宝坊公开报道确认品类", "________________________________"],
            ["6", "大名五百居香肠", "公开存在 500 g 袋及礼盒等多规格", "________________________________"],
        ],
        [0.48, 1.55, 2.25, 2.55],
        7.8,
    )
    add_heading(doc, "二、每个 SKU 必须提供的结构数据", 1)
    add_bullets(doc, [
        "最大外廓：长×宽×高（瓶/罐可填最大直径×高度），包含瓶盖、封口、提手、鼓包和其他凸起；单位 mm。",
        "抽取三件样品分别实测并填写最大值；软袋在正常充填、自然放置状态测量，不用净含量推算。",
        "单件毛重、计划装盒数量、正常摆放方向、不可倒置/不可受压区、易碎、防潮、避光、冷链或异味隔离要求。",
        "六面高清照片（正、背、左、右、顶、底），每张建议长边不低于 3000 px；另附允许用于展板的透明底 PNG 或原始商品图。",
        "真实标签全文、条码位置、日期/批号位置，以及必须保留的最小字号、可视面积和法务/质量审核要求。",
    ])
    add_status_box(doc, "特别说明", "净含量（mL/g/kg）不是包装毛重，也不能推导外廓尺寸。资料未回填前，现有 60×60×35 mm 模型仅为占位件。")

    doc.add_page_break()
    add_heading(doc, "三、SKU 数据回填表（1—3）", 1)
    for idx in range(1, 4):
        add_heading(doc, f"SKU {idx}", 2)
        table(
            doc,
            ["字段", "品牌方填写", "字段", "品牌方填写"],
            [
                ["商品全称", "________________", "条码", "________________"],
                ["净含量", "________________", "单件毛重/g", "________________"],
                ["长×宽×高/mm", "________________", "三件最大值/mm", "________________"],
                ["计划数量", "________________", "正常朝向", "________________"],
                ["包装材质", "________________", "易碎等级", "□高 □中 □低"],
                ["储存/隔离", "________________", "不可受压区", "________________"],
                ["六面图文件名", "________________", "样品编号", "________________"],
            ],
            [1.02, 2.28, 1.08, 2.4],
            7.8,
        )
        doc.add_paragraph().paragraph_format.space_after = Pt(1)

    doc.add_page_break()
    add_heading(doc, "四、SKU 数据回填表（4—6）", 1)
    for idx in range(4, 7):
        add_heading(doc, f"SKU {idx}", 2)
        table(
            doc,
            ["字段", "品牌方填写", "字段", "品牌方填写"],
            [
                ["商品全称", "________________", "条码", "________________"],
                ["净含量", "________________", "单件毛重/g", "________________"],
                ["长×宽×高/mm", "________________", "三件最大值/mm", "________________"],
                ["计划数量", "________________", "正常朝向", "________________"],
                ["包装材质", "________________", "易碎等级", "□高 □中 □低"],
                ["储存/隔离", "________________", "不可受压区", "________________"],
                ["六面图文件名", "________________", "样品编号", "________________"],
            ],
            [1.02, 2.28, 1.08, 2.4],
            7.8,
        )
        doc.add_paragraph().paragraph_format.space_after = Pt(1)

    doc.add_page_break()
    add_heading(doc, "五、品牌资产与图片交付清单", 1)
    table(
        doc,
        ["资料", "要求", "品牌方填写/勾选"],
        [
            ["品牌 Logo", "AI/SVG/PDF 矢量文件；最小尺寸、保护区、禁用规范", "□已附  文件名：____________"],
            ["品牌色与字体", "CMYK/RGB/Pantone；字体名称及可用许可", "□已附  □无指定"],
            ["商品图片", "六面图、透明底图、场景图；标注摄影/版权归属", "□已附  共____张"],
            ["法定文案", "品名、配料、过敏原、净含量、储存、生产者、许可/标准", "□已附  □需最终审核"],
            ["样品", "每个 SKU 建议 3 件，标注批次；供测量与试装", "□可提供  □不可提供"],
            ["审核联系人", "姓名、职务、手机/邮箱；能确认品牌与标签资料", "________________________"],
        ],
        [1.25, 3.75, 1.78],
        8.2,
    )
    add_heading(doc, "六、拟申请的有限授权范围", 1)
    add_body(doc, "若品牌方同意，请在下一页明确授权主体、素材清单、期限和发布边界。建议的最小范围为：允许申请方在 2026“和氏璧杯”参赛文件、现场/线上评审展示和答辩 PPT 中使用指定的邯宝坊名称、Logo、商品名称、商品图片和经确认的标签信息；允许为参赛目的制作不销售的包装样机；不包含商品生产、销售、广告代言、商标注册、再授权或对品牌方构成合作/获奖背书。")
    add_bullets(doc, [
        "公开作品集、社交媒体和展览是否允许，应由品牌方单独勾选；未勾选视为不允许。",
        "品牌方保留对最终品牌呈现、商品事实和法定食品文案的审核权；结构创意和参赛作品著作权按赛事规则及作者实际权属处理。",
        "商品停产、换包装或资料变更时，品牌方应提示申请方停止使用旧素材；申请方不据此宣称官方联名、推荐或量产。",
    ])

    doc.add_page_break()
    add_heading(doc, "七、品牌方授权确认（须由有权代表填写）", 1)
    add_status_box(doc, "签署提示", "本页空白时视为未授权。电子回传件应包含签署人、日期与公章/合同章；若使用电子签章，请附可验证文件。")
    add_body(doc, "授权方全称：________________________________________________________")
    add_body(doc, "统一社会信用代码：__________________________________________________")
    add_body(doc, "权利基础：□商标权人  □著作权人  □合法被许可人  □其他：____________")
    add_body(doc, "经办人/职务：________________________  联系方式：____________________")
    add_body(doc, "授权素材清单（请写文件名或附件编号）：")
    add_body(doc, "____________________________________________________________________")
    add_body(doc, "授权用途（可多选）：")
    add_body(doc, "□赛事投稿文件  □线上/现场评审  □答辩 PPT  □不销售样机制作")
    add_body(doc, "□赛后作品集（仅展示设计）  □社交媒体展示  □展览  □其他：____________")
    add_body(doc, "授权地域：__________________  授权期限：自________至________")
    add_body(doc, "审核要求：□使用前逐稿确认  □仅核对事实/标签  □按所附规范使用即可")
    add_body(doc, "其他限制：__________________________________________________________")
    doc.add_paragraph()
    table(
        doc,
        ["授权方签署", "申请方确认"],
        [
            ["有权代表签字：________________\n\n公章/合同章：\n\n日期：______年____月____日", "申请人签字：________________\n\n院校/单位（如适用）：____________\n\n日期：______年____月____日"],
        ],
        [3.35, 3.35],
        9,
    )
    add_heading(doc, "附件建议", 1)
    add_bullets(doc, [
        "A1：《解围》6 张 A4 设计展板（现阶段商品为中性占位，供理解结构，不作品牌终稿）。",
        "A2：闭合、解锁、开启三状态 CAD/STEP 预览及刀模说明。",
        "A3：《邯宝坊商品公开证据档案》与商品证据 CSV。",
        "A4：品牌方回填的 SKU 表、六面照片、标签文案、素材清单和样品交接记录。",
    ])
    add_body(doc, "资料接收后，设计方预计在 3—5 个工作日内完成参数重算和数字样机更新；实物白样、满载开合与跌落/振动测试时间另计。该时限是项目计划，不是赛事或品牌方承诺。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    normalize_docx_archive(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
